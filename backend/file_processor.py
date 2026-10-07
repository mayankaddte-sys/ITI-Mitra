import os
from typing import Optional
import PyPDF2
from docx import Document
from dotenv import load_dotenv

load_dotenv()

UPLOAD_DIR = os.getenv("UPLOAD_DIR", "./uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

class FileProcessor:
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        text = ""
        try:
            with open(file_path, "rb") as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
        except Exception as e:
            raise Exception(f"Error reading PDF: {str(e)}")
        return text

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        text = ""
        try:
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                text += "\n"
        except Exception as e:
            raise Exception(f"Error reading DOCX: {str(e)}")
        return text

    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        try:
            with open(file_path, "r", encoding="utf-8") as txt_file:
                return txt_file.read()
        except Exception as e:
            raise Exception(f"Error reading TXT: {str(e)}")

    @staticmethod
    def process_file(file_path: str) -> str:
        _, file_ext = os.path.splitext(file_path)
        file_ext = file_ext.lower()

        if file_ext == ".pdf":
            return FileProcessor.extract_text_from_pdf(file_path)
        elif file_ext == ".docx":
            return FileProcessor.extract_text_from_docx(file_path)
        elif file_ext == ".txt":
            return FileProcessor.extract_text_from_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_ext}")
